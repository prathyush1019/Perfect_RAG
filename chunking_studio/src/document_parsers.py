import io
import os
from typing import Union, Tuple


def parse_uploaded_file(file_obj, filename: str) -> Tuple[str, str]:
    """
    Parses uploaded file or file-like object and extracts text.
    Returns: (extracted_text, file_type)
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext in [".txt", ".md", ".markdown", ".csv", ".json"]:
        if hasattr(file_obj, "read"):
            content = file_obj.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="replace"), ext
            return str(content), ext
        elif isinstance(file_obj, str):
            with open(file_obj, "r", encoding="utf-8", errors="replace") as f:
                return f.read(), ext
        elif isinstance(file_obj, bytes):
            return file_obj.decode("utf-8", errors="replace"), ext

    elif ext == ".pdf":
        try:
            import pypdf
            if hasattr(file_obj, "read"):
                # Seek to 0 if possible
                if hasattr(file_obj, "seek"):
                    file_obj.seek(0)
                reader = pypdf.PdfReader(file_obj)
            else:
                reader = pypdf.PdfReader(file_obj)

            text_pages = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                text_pages.append(page_text)
            return "\n\n".join(text_pages), "pdf"
        except Exception as e:
            raise RuntimeError(f"Error reading PDF file: {str(e)}")

    elif ext in [".docx", ".doc"]:
        try:
            import docx
            if hasattr(file_obj, "read"):
                if hasattr(file_obj, "seek"):
                    file_obj.seek(0)
                doc = docx.Document(file_obj)
            else:
                doc = docx.Document(file_obj)

            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract tables
            table_texts = []
            for t in doc.tables:
                rows_text = []
                for row in t.rows:
                    cols = [cell.text.strip() for cell in row.cells]
                    rows_text.append("| " + " | ".join(cols) + " |")
                if rows_text:
                    table_texts.append("\n".join(rows_text))

            full_text = "\n\n".join(paragraphs)
            if table_texts:
                full_text += "\n\n" + "\n\n".join(table_texts)
            return full_text, "docx"
        except Exception as e:
            raise RuntimeError(f"Error reading DOCX file: {str(e)}")

    else:
        # Default text fallback
        if hasattr(file_obj, "read"):
            content = file_obj.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="replace"), ext
            return str(content), ext
        return str(file_obj), ext
